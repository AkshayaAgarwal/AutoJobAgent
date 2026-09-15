#!/usr/bin/env python3
"""
Job Data Cleaner - Apify Seek/LinkedIn Job JSON Cleaner
Extracts critical job info, contacts, requirements, and main points
Outputs clean XLSX ready for cover letter generation and applications
"""

import json
import re
from html.parser import HTMLParser
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import sys

class HTMLStripper(HTMLParser):
    """Remove HTML tags from text"""
    def __init__(self):
        super().__init__()
        self.reset()
        self.strict = False
        self.convert_charrefs = True
        self.text = []
        
    def handle_data(self, data):
        self.text.append(data)
        
    def get_text(self):
        return ''.join(self.text)

def strip_html(html_text):
    """Strip HTML tags and clean whitespace"""
    if not html_text:
        return ""
    stripper = HTMLStripper()
    try:
        stripper.feed(html_text)
        text = stripper.get_text()
        # Clean multiple spaces and newlines
        text = re.sub(r'\s+', ' ', text).strip()
        return text
    except:
        return html_text

def extract_requirements(content_text):
    """Extract essential and nice-to-have requirements"""
    essential = []
    nice_to_have = []
    
    # Split by sections
    sections = content_text.split("Nice to have")
    
    if len(sections) > 0:
        # Extract essential requirements
        essential_match = re.search(r'Essential\s*(.*?)(?:Nice|$)', sections[0], re.DOTALL)
        if essential_match:
            essential_text = essential_match.group(1)
            # Extract bullet points
            bullets = re.findall(r'[-•]\s*(.+?)(?=[-•]|$)', essential_text)
            essential = [b.strip() for b in bullets if b.strip()]
    
    if len(sections) > 1:
        # Extract nice-to-have requirements
        bullets = re.findall(r'[-•]\s*(.+?)(?=[-•]|$)', sections[1])
        nice_to_have = [b.strip() for b in bullets if b.strip()]
    
    return essential, nice_to_have

def extract_responsibilities(content_text):
    """Extract main responsibilities from job description"""
    responsibilities = []
    
    # Look for "What you'll actually do" or similar sections
    match = re.search(r'What.*?do(.*?)(?:What we need|$)', content_text, re.DOTALL | re.IGNORECASE)
    if match:
        resp_text = match.group(1)
        bullets = re.findall(r'[-•]\s*(.+?)(?=[-•]|$)', resp_text)
        responsibilities = [r.strip() for r in bullets if r.strip() and len(r.strip()) > 10][:8]  # Top 8
    
    return responsibilities

def extract_main_points(job_data):
    """Extract key main points for cover letter reference"""
    points = {
        'job_hook': job_data.get('content', {}).get('jobHook', ''),
        'company_focus': '',
        'key_tech_stack': [],
        'key_responsibilities': [],
        'company_benefits': []
    }
    
    # Get company focus from unEditedContent
    unedited = job_data.get('content', {}).get('unEditedContent', '')
    unedited_clean = strip_html(unedited)
    
    # Extract company focus
    about_match = re.search(r'About us\s*(.*?)(?:About the role|$)', unedited_clean, re.IGNORECASE)
    if about_match:
        points['company_focus'] = about_match.group(1).strip()[:200]
    
    # Extract tech stack
    tech_patterns = [
        r'React', r'Next\.js', r'TypeScript', r'JavaScript', r'Python', r'C#', r'\.NET',
        r'Node\.js?', r'SQL', r'AWS', r'Azure', r'LLM', r'Claude', r'GPT'
    ]
    for tech in tech_patterns:
        if re.search(tech, unedited_clean, re.IGNORECASE):
            points['key_tech_stack'].append(tech.replace(r'\.', '.').replace('?', ''))
    
    # Extract key responsibilities
    points['key_responsibilities'] = extract_responsibilities(unedited_clean)[:5]
    
    # Extract benefits
    benefits_match = re.search(r'What you get(.*?)(?:How to apply|$)', unedited_clean, re.IGNORECASE)
    if benefits_match:
        benefits_text = benefits_match.group(1)
        bullets = re.findall(r'[-•]\s*(.+?)(?=[-•]|$)', benefits_text)
        points['company_benefits'] = [b.strip() for b in bullets if b.strip()][:5]
    
    return points

def clean_salary(salary_str):
    """Clean and standardize salary"""
    if not salary_str:
        return "Not specified"
    return salary_str.strip()

def process_job_data(json_data):
    """Process single or multiple job records"""
    # Handle both single job object and array of jobs
    if isinstance(json_data, dict):
        jobs = [json_data]
    elif isinstance(json_data, list):
        jobs = json_data
    else:
        raise ValueError("JSON must be a job object or array of jobs")
    
    cleaned_jobs = []
    
    for job in jobs:
        try:
            unedited_content = job.get('content', {}).get('unEditedContent', '')
            clean_content = strip_html(unedited_content)
            
            essential_req, nice_req = extract_requirements(clean_content)
            main_points = extract_main_points(job)
            
            cleaned_job = {
                'job_id': job.get('id', 'N/A'),
                'job_title': job.get('title', 'N/A'),
                'company_name': job.get('advertiser', {}).get('name', 'N/A'),
                'location': job.get('joblocationInfo', {}).get('displayLocation', 'N/A'),
                'salary': clean_salary(job.get('salary', '')),
                'work_type': job.get('workTypes', 'N/A'),
                'work_arrangement': job.get('workArrangements', 'N/A'),
                'contact_email': ', '.join(job.get('emails', ['Not specified'])),
                'apply_link': job.get('applyLink', ''),
                'job_link': job.get('jobLink', ''),
                'company_website': job.get('companyProfile', {}).get('website', ''),
                'company_size': job.get('companyProfile', {}).get('size', 'N/A'),
                'industry': job.get('classificationInfo', {}).get('subClassification', 'N/A'),
                'job_hook': main_points['job_hook'],
                'company_focus': main_points['company_focus'],
                'tech_stack': ', '.join(main_points['key_tech_stack']),
                'key_responsibilities': '\n'.join(main_points['key_responsibilities']),
                'company_benefits': '\n'.join(main_points['company_benefits']),
                'essential_requirements': '\n'.join(essential_req),
                'nice_to_have': '\n'.join(nice_req),
                'num_applicants': job.get('numApplicants', 'N/A'),
                'listed_date': job.get('listedAt', 'N/A')[:10]  # Date only
            }
            cleaned_jobs.append(cleaned_job)
        except Exception as e:
            print(f"Error processing job {job.get('id', 'Unknown')}: {str(e)}")
            continue
    
    return cleaned_jobs

def create_xlsx(cleaned_jobs, output_path):
    """Create professional XLSX workbook with cleaned job data"""
    wb = Workbook()
    
    # Create main job data sheet
    ws_jobs = wb.active
    ws_jobs.title = "Job Opportunities"
    
    # Define headers for main sheet
    headers = [
        'Job ID',
        'Job Title',
        'Company',
        'Location',
        'Salary',
        'Work Type',
        'Arrangement',
        'Contact Email',
        'Website',
        'Company Size',
        'Industry',
        'Job Link',
        'Apply Link',
        'Applicants',
        'Posted Date'
    ]
    
    # Write headers with formatting
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws_jobs.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Write job data
    for row_idx, job in enumerate(cleaned_jobs, 2):
        ws_jobs.cell(row=row_idx, column=1, value=job['job_id'])
        ws_jobs.cell(row=row_idx, column=2, value=job['job_title'])
        ws_jobs.cell(row=row_idx, column=3, value=job['company_name'])
        ws_jobs.cell(row=row_idx, column=4, value=job['location'])
        ws_jobs.cell(row=row_idx, column=5, value=job['salary'])
        ws_jobs.cell(row=row_idx, column=6, value=job['work_type'])
        ws_jobs.cell(row=row_idx, column=7, value=job['work_arrangement'])
        ws_jobs.cell(row=row_idx, column=8, value=job['contact_email'])
        ws_jobs.cell(row=row_idx, column=9, value=job['company_website'])
        ws_jobs.cell(row=row_idx, column=10, value=job['company_size'])
        ws_jobs.cell(row=row_idx, column=11, value=job['industry'])
        ws_jobs.cell(row=row_idx, column=12, value=job['job_link'])
        ws_jobs.cell(row=row_idx, column=13, value=job['apply_link'])
        ws_jobs.cell(row=row_idx, column=14, value=job['num_applicants'])
        ws_jobs.cell(row=row_idx, column=15, value=job['listed_date'])
        
        # Add borders and alignment
        for col in range(1, len(headers) + 1):
            cell = ws_jobs.cell(row=row_idx, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    
    # Set column widths
    ws_jobs.column_dimensions['A'].width = 12
    ws_jobs.column_dimensions['B'].width = 25
    ws_jobs.column_dimensions['C'].width = 20
    ws_jobs.column_dimensions['D'].width = 18
    ws_jobs.column_dimensions['E'].width = 18
    ws_jobs.column_dimensions['F'].width = 12
    ws_jobs.column_dimensions['G'].width = 14
    ws_jobs.column_dimensions['H'].width = 22
    ws_jobs.column_dimensions['I'].width = 20
    ws_jobs.column_dimensions['J'].width = 16
    ws_jobs.column_dimensions['K'].width = 18
    ws_jobs.column_dimensions['L'].width = 18
    ws_jobs.column_dimensions['M'].width = 18
    ws_jobs.column_dimensions['N'].width = 12
    ws_jobs.column_dimensions['O'].width = 12
    
    # Create detailed job info sheet
    ws_details = wb.create_sheet("Job Details")
    detail_headers = ['Job ID', 'Job Title', 'Company', 'Job Hook', 'Company Focus', 
                      'Tech Stack', 'Key Responsibilities', 'Company Benefits', 
                      'Essential Requirements', 'Nice-to-Have']
    
    for col, header in enumerate(detail_headers, 1):
        cell = ws_details.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    for row_idx, job in enumerate(cleaned_jobs, 2):
        ws_details.cell(row=row_idx, column=1, value=job['job_id'])
        ws_details.cell(row=row_idx, column=2, value=job['job_title'])
        ws_details.cell(row=row_idx, column=3, value=job['company_name'])
        ws_details.cell(row=row_idx, column=4, value=job['job_hook'])
        ws_details.cell(row=row_idx, column=5, value=job['company_focus'])
        ws_details.cell(row=row_idx, column=6, value=job['tech_stack'])
        ws_details.cell(row=row_idx, column=7, value=job['key_responsibilities'])
        ws_details.cell(row=row_idx, column=8, value=job['company_benefits'])
        ws_details.cell(row=row_idx, column=9, value=job['essential_requirements'])
        ws_details.cell(row=row_idx, column=10, value=job['nice_to_have'])
        
        for col in range(1, len(detail_headers) + 1):
            cell = ws_details.cell(row=row_idx, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    
    # Set column widths for details sheet
    ws_details.column_dimensions['A'].width = 12
    ws_details.column_dimensions['B'].width = 25
    ws_details.column_dimensions['C'].width = 20
    ws_details.column_dimensions['D'].width = 30
    ws_details.column_dimensions['E'].width = 30
    ws_details.column_dimensions['F'].width = 25
    ws_details.column_dimensions['G'].width = 35
    ws_details.column_dimensions['H'].width = 30
    ws_details.column_dimensions['I'].width = 35
    ws_details.column_dimensions['J'].width = 30
    
    # Save workbook
    wb.save(output_path)
    print(f"✓ XLSX file created successfully: {output_path}")
    print(f"✓ Total jobs processed: {len(cleaned_jobs)}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python job_data_cleaner.py <input_json_file> [output_xlsx_file]")
        print("\nExample: python job_data_cleaner.py jobs.json jobs_cleaned.xlsx")
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "jobs_cleaned.xlsx"
    
    # Load JSON data
    try:
        with open(input_file, 'r', encoding='utf-8') as f:
            json_data = json.load(f)
            print(f"✓ Loaded JSON from {input_file}")
    except FileNotFoundError:
        print(f"Error: File '{input_file}' not found")
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON format - {str(e)}")
        sys.exit(1)
    
    # Process job data
    try:
        cleaned_jobs = process_job_data(json_data)
        print(f"✓ Processed {len(cleaned_jobs)} jobs")
        
        if len(cleaned_jobs) == 0:
            print("Warning: No jobs were processed successfully")
            sys.exit(1)
        
        # Create XLSX
        create_xlsx(cleaned_jobs, output_file)
        print(f"\n✅ Job data cleaned and exported to: {output_file}")
        
    except Exception as e:
        print(f"Error processing data: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    main()